# app.py - Gerador de Declaração de Margem (versão final com detalhes de consignados)
# --- Página inicial ---

import streamlit as st
import gspread
from oauth2client.service_account import ServiceAccountCredentials
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from io import BytesIO
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from num2words import num2words
import unicodedata
import re
import streamlit as st

# ---------- SESSÃO ----------
if "logado" not in st.session_state:
    st.session_state.logado = False
if "usuario" not in st.session_state:
    st.session_state.usuario = None

 


# <<< ADIÇÃO — FUNÇÃO PARA FORÇAR ARIAL 12 EM TODO DOCUMENTO >>>
def aplicar_fonte_arial_12(documento):
    for p in documento.paragraphs:
        for run in p.runs:
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(12)

    for table in documento.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Arial"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
                        run.font.size = Pt(12)
# <<< FIM DA ADIÇÃO >>>


# --- Configuração da página ---
st.set_page_config(page_title="Gerador de Declaração de Margem", page_icon="💼", layout="centered")

# --- Estilos personalizados ---
st.markdown(
    """
    <style>
    .stApp { background-color: #FFE992 !important; }
    .stButton>button {
        background-color: #FFF4BC !important;
        color: black !important;
        border: 1px solid #d4be6a !important;
        padding: 0.6rem 1.2rem !important;
        border-radius: 8px !important;
        font-weight: bold !important;
    }
    .stButton>button:hover {
        background-color: #ffe27a !important;
        border-color: #c6a94d !important;
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.title("💼 Gera Margem")
st.write("Selecione um nome para gerar automaticamente a declaração.")


# --- Helpers ---
def normalize_header(s: str) -> str:
    if s is None:
        return ""
    s = str(s).strip().upper()
    s = unicodedata.normalize("NFKD", s).encode("ASCII", "ignore").decode("ASCII")
    s = " ".join(s.split())
    return s


def parse_decimal(value) -> Decimal:
    if value is None:
        return Decimal("0.00")
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if isinstance(value, int):
        return Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if isinstance(value, float):
        return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    s = str(value).strip()
    if s == "":
        return Decimal("0.00")

    s = s.replace("R$", "").replace("r$", "")
    s = s.replace("\u00A0", "").replace(" ", "")

    if "." in s and "," in s:
        last_dot = s.rfind(".")
        last_comma = s.rfind(",")
        if last_comma > last_dot:
            s = s.replace(".", "")
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        if "," in s and "." not in s:
            s = s.replace(",", ".")

    filtered = "".join(ch for ch in s if ch.isdigit() or ch in ".-+")
    if filtered in ("", ".", "-", "+"):
        return Decimal("0.00")
    try:
        d = Decimal(filtered)
    except InvalidOperation:
        return Decimal("0.00")
    return d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def extenso_brl(valor: Decimal) -> str:
    valor = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    sinal = ""
    if valor < 0:
        sinal = "menos "
        valor = abs(valor)
    inteiro = int(valor)
    centavos = int((valor - Decimal(inteiro)) * 100)
    if inteiro == 0 and centavos == 0:
        return "zero reais"
    if centavos:
        texto = f"{num2words(inteiro, lang='pt_BR')} reais e {num2words(centavos, lang='pt_BR')} centavos"
    else:
        texto = f"{num2words(inteiro, lang='pt_BR')} reais"
    texto = re.sub(r"\s+", " ", texto).strip()
    return f"{sinal}{texto}"


def format_brl(valor: Decimal) -> str:
    q = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    s = f"{q:,.2f}"
    s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


# --- ADIÇÃO — permitir negrito com <b> ... </b> ---
def _create_runs_with_bold_tags(paragraph, text_with_tags: str):
    parts = re.split(r'(<b>.*?<\/b>)', text_with_tags)
    for part in parts:
        if not part:
            continue
        m = re.match(r'^<b>(.*?)</b>$', part)
        if m:
            run = paragraph.add_run(m.group(1))
            run.bold = True
        else:
            paragraph.add_run(part)


def replace_in_doc(doc: Document, subs: dict):
    for p in doc.paragraphs:
        if not p.text:
            continue
        p_text = p.text
        replaced_any = False
        for chave, valor in subs.items():
            if chave in p_text:
                p_text = p_text.replace(chave, str(valor))
                replaced_any = True
        if replaced_any and p_text != p.text:
            for _ in range(len(p.runs)):
                p.runs[0]._element.getparent().remove(p.runs[0]._element)
            _create_runs_with_bold_tags(p, p_text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text
                replaced_any = False
                for chave, valor in subs.items():
                    if chave in cell_text:
                        cell_text = cell_text.replace(chave, str(valor))
                        replaced_any = True
                if replaced_any:
                    cell._tc.clear_content()
                    p_new = cell.add_paragraph()
                    _create_runs_with_bold_tags(p_new, cell_text)


def ajustar_texto_sem_consignados(doc: Document):
    """
    Remove APENAS o trecho:
    ', com margem comprometida no valor de ..., restando uma margem livre de ...'
    mantendo o resto do texto intacto e SEM duplicar parágrafos.
    """

    padrao = re.compile(
        r",\s*com margem comprometida no valor de.*?(\.)",
        flags=re.IGNORECASE
    )

    for p in doc.paragraphs:
        if "Gerando uma MARGEM CONSIGNÁVEL" in p.text:
            novo_texto = padrao.sub(".", p.text)

            # limpa runs antigos
            for _ in range(len(p.runs)):
                p.runs[0]._element.getparent().remove(p.runs[0]._element)

            p.add_run(novo_texto)
            break





# --- Conexão com Google Sheets ---
scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
creds = ServiceAccountCredentials.from_json_keyfile_name("credentials.json", scope)
client = gspread.authorize(creds)

sheet = client.open_by_url(
    "https://docs.google.com/spreadsheets/d/1jRpDiEc9kaEDNAGjjE-afu1pS9S9xk31L5kzxMoTxo4/edit?gid=0"
).sheet1

raw_records = sheet.get_all_records(head=2)

if not raw_records:
    st.error("A planilha não retornou registros. Verifique o head/linha de cabeçalho.")
    st.stop()

header_map = {}
header_row = sheet.row_values(2)
for h in header_row:
    header_map[normalize_header(h)] = h


def get_field(record, desired_name):
    key = header_map.get(normalize_header(desired_name))
    if key and key in record:
        return record.get(key)
    for k in record.keys():
        if normalize_header(k) == normalize_header(desired_name):
            return record.get(k)
    return None


nomes = [get_field(r, "NOME") for r in raw_records]
nomes = [n for n in nomes if n is not None and str(n).strip() != ""]

nome_selecionado = st.selectbox("Selecione o nome:", nomes)

if nome_selecionado:
    pessoa = next((r for r in raw_records if str(get_field(r, "NOME")).strip() == str(nome_selecionado).strip()), None)

    if not pessoa:
        st.error("Não foi possível localizar a pessoa selecionada.")
    else:
        raw_salario = get_field(pessoa, "SALÁRIO") or get_field(pessoa, "REMUNERAÇÃO") or get_field(pessoa, "SALARIO")
        salario = parse_decimal(raw_salario)

        raw_margem30 = get_field(pessoa, "MARGEM 30%") or get_field(pessoa, "MARGEM30%") or get_field(pessoa, "MARGEM30")
        if raw_margem30 not in (None, ""):
            margem_total = parse_decimal(raw_margem30)
        else:
            margem_total = (salario * Decimal("0.30")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

        vinculo_raw = get_field(pessoa, "VÍNCULO") or get_field(pessoa, "VINCULO") or ""
        vinculo = str(vinculo_raw).strip()

        # <<< ADIÇÃO TRATAMENTO SR./SRA. >>>
        sexo_raw = get_field(pessoa, "SEXO") or get_field(pessoa, "GENERO") or ""
        sexo = str(sexo_raw).strip().upper()

        if sexo.startswith("F"):
            tratamento = "a Sra."
        elif sexo.startswith("M"):
            tratamento = "o Sr."
        else:
            tratamento = "a Sra."
        # <<< FIM DA ADIÇÃO >>>

        vinculo_upper = vinculo.upper()
        if "APOSEN" in vinculo_upper:
            vinculo_doc = "APOSENTADA"
        elif "PENSION" in vinculo_upper:
            vinculo_doc = "PENSIONISTA"
        else:
            vinculo_doc = vinculo if vinculo else "---"

        matricula = get_field(pessoa, "MATRÍCULA") or get_field(pessoa, "MATRICULA") or "---"
        cpf = get_field(pessoa, "CPF") or "---"

        consignados_vals = []
        possiveis = [
            "CONSIGNADO 1","CONSIGNADO1",
            "CONSIGNADO 2","CONSIGNADO2",
            "CONSIGNADO 3","CONSIGNADO3",
            "CONSIGNADO 4","CONSIGNADO4",
            "CONSIGNADO 5","CONSIGNADO5",
            "EMPRÉSTIMO 1","EMPRESTIMO1",
            "EMPRÉSTIMO 2","EMPRESTIMO2",
            "EMPRÉSTIMO 3","EMPRESTIMO3",
            "EMPRÉSTIMO 4","EMPRESTIMO4",
            "EMPRÉSTIMO 5","EMPRESTIMO5"
        ]
        for nome_col in possiveis:
            v = get_field(pessoa, nome_col)
            if v is None or (isinstance(v, str) and v.strip() == ""):
                continue
            d = parse_decimal(v)
            if d != Decimal("0.00"):
                consignados_vals.append(d)

        if not consignados_vals:
            single = get_field(pessoa, "CONSIGNADO") or get_field(pessoa, "EMPRESTIMO")
            if single not in (None, ""):
                d = parse_decimal(single)
                if d != Decimal("0.00"):
                    consignados_vals.append(d)

        margem_comprometida = sum(consignados_vals, Decimal("0.00")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        margem_livre = (margem_total - margem_comprometida).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

        consignado_placeholders = {}
        for i in range(1, 6):
            if i <= len(consignados_vals):
                val = consignados_vals[i-1]
                num_full = format_brl(val)
                cons_text = f"<b>{num_full}</b>"
                consignado_placeholders[f"{{{{CONSIGNADO{i}_NUM}}}}"] = cons_text
                consignado_placeholders[f"{{{{CONSIGNADO{i}_EXT}}}}"] = extenso_brl(val)
            else:
                consignado_placeholders[f"{{{{CONSIGNADO{i}_NUM}}}}"] = ""
                consignado_placeholders[f"{{{{CONSIGNADO{i}_EXT}}}}"] = ""

        if consignados_vals:
            partes = []
            for v in consignados_vals:
                num_full = format_brl(v)
                ext_text = extenso_brl(v)
                partes.append(f"<b>{num_full}</b> ({ext_text})")
            consignados_lista_text = ", ".join(partes)
        else:
            consignados_lista_text = ""

        st.subheader("📊 Resumo (confira os valores)")
        st.write(f"**Nome:** {nome_selecionado}")
        st.write(f"**Tratamento:** {tratamento}")
        st.write(f"**Vínculo:** {vinculo_doc}")
        st.write(f"**Matrícula:** {matricula}")
        st.write(f"**CPF:** {cpf}")
        st.write(f"**Salário (num):** {format_brl(salario)}")
        st.write(f"**Salário (extenso):** {extenso_brl(salario)}")
        st.write(f"**Margem Total (30%):** {format_brl(margem_total)} ({extenso_brl(margem_total)})")
        st.write(f"**Margem Comprometida:** {format_brl(margem_comprometida)} ({extenso_brl(margem_comprometida)})")
        st.write(f"**Margem Livre:** {format_brl(margem_livre)} ({extenso_brl(margem_livre)})")

        if margem_comprometida > margem_total:
            st.warning("⚠️ Margem comprometida maior que a margem total.")

        if st.button("📄 Gerar Declaração"):
            modelo = "DECLARAÇÃO_DE_MARGEM_MODELO.docx"
            try:
                doc = Document(modelo)
            except Exception as e:
                st.error(f"Não foi possível abrir o modelo '{modelo}': {e}")
                st.stop()

            if not consignados_vals:
                ajustar_texto_sem_consignados(doc)


            meses_pt = {
                1: "janeiro",2: "fevereiro",3: "março",4: "abril",5: "maio",6: "junho",
                7: "julho",8: "agosto",9: "setembro",10: "outubro",11: "novembro",12: "dezembro"
            }
            hoje = datetime.now()
            data_pt = f"{hoje.day} de {meses_pt[hoje.month]} de {hoje.year}"

            substituicoes = {
                "{{TRATAMENTO}}": tratamento,   # <<< AQUI FOI ADICIONADO >>>
                "{{NOME}}": f"<b>{nome_selecionado}</b>",
                "{{CPF}}": f"<b>{cpf}</b>",
                "{{MATRICULA}}": f"<b>{matricula}</b>",
                "{{VINCULO}}": f"<b>{vinculo_doc}</b>",
                "{{SALARIO_NUM}}": f"<b>{format_brl(salario)}</b>",
                "{{SALARIO_EXT}}": extenso_brl(salario),
                "{{MARGEM_TOTAL_NUM}}": f"<b>{format_brl(margem_total)}</b>",
                "{{MARGEM_TOTAL_EXT}}": extenso_brl(margem_total),
                "{{CONSIGNADOS_LISTA}}": consignados_lista_text,
                "{{MARGEM_LIVRE_NUM}}": f"<b>{format_brl(margem_livre)}</b>",
                "{{MARGEM_LIVRE_EXT}}": extenso_brl(margem_livre),
                "{{DATA}}": data_pt
            }

            substituicoes.update(consignado_placeholders)

            replace_in_doc(doc, substituicoes)

            aplicar_fonte_arial_12(doc)

            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            safe_name = re.sub(r"[^A-Za-z0-9 _-]", "", str(nome_selecionado))
            file_name = f"Declaracao_{safe_name}.docx"

            st.success("✅ Declaração gerada com sucesso!")
            st.download_button(
                "⬇️ Baixar declaração",
                data=buffer,
                file_name=file_name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
